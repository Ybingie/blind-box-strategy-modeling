"""Read the supplied workbook and prepare the model inputs without editing Excel."""
from collections import defaultdict
import csv
from datetime import datetime
import json
from pathlib import Path
from statistics import median

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent


def prepare_data():
    workbook = load_workbook(ROOT / 'DEATH_NOTE_My_Dear_Light.xlsx', read_only=True, data_only=True)
    try:
        role_rows = workbook['02_Character_Probabilities'].iter_rows(min_row=3, values_only=True)
        roles = {row[1]: row[0] for row in role_rows if isinstance(row[0], int)}
        source_rows = [row for row in workbook['03_Secondary_Market_Raw_Data'].iter_rows(values_only=True)
                       if isinstance(row[0], int)]
        records = []
        groups = defaultdict(list)
        for row in source_rows:
            if row[19] != 'Yes' or not isinstance(row[12], (int, float)):
                continue
            date = row[5].strftime('%Y-%m-%d') if isinstance(row[5], datetime) else str(row[5])
            record = {'record_id': row[0], 'role_index': roles[row[1]], 'role': row[1],
                      'collection_date': date, 'unit_cost': float(row[12]), 'evidence_id': row[21]}
            records.append(record)
            groups[(roles[row[1]], row[4])].append(record)
        flagged_ids = {
            record['evidence_id'] for group in groups.values() if len(group) > 1 for record in group
        }
        date_ids = [
            record['evidence_id'] for record in records if not record['collection_date'].startswith('2026-')
        ]
        flagged_ids.update(date_ids)
        retained = [record for record in records if record['evidence_id'] not in flagged_ids]

        def medians(items):
            return [median([r['unit_cost'] for r in items if r['role_index'] == i]) for i in range(1, 9)]

        audit = {'source_records': len(source_rows), 'usable_records': len(records),
                 'unique_evidence_ids': len({r['evidence_id'] for r in records}),
                 'unique_listing_urls': len({row[4] for row in source_rows
                                            if row[19] == 'Yes' and isinstance(row[12], (int, float))}),
                 'same_design_repeated_groups': sum(len(group) > 1 for group in groups.values()),
                 'unresolved_date_ids': date_ids,
                 'conservative_check_excluded_ids': sorted(flagged_ids),
                 'conservative_check_records': len(retained),
                 'baseline_medians': medians(records), 'conservative_check_medians': medians(retained),
                 'limitations': ['Convenience sample of listings, not completed transactions.',
                                 'Search terms and seller identities were not systematically recorded.',
                                 'Recorded screenshot names do not verify archived screenshots.']}
        data = ROOT / 'data'
        data.mkdir(exist_ok=True)
        with (data / 'market_prices_cleaned.csv').open('w', encoding='utf-8-sig', newline='') as stream:
            writer = csv.DictWriter(stream, fieldnames=list(records[0]))
            writer.writeheader()
            writer.writerows(records)
        (data / 'market_trace_audit.json').write_text(json.dumps(audit, indent=2), encoding='utf-8')
        return audit
    finally:
        workbook.close()


if __name__ == '__main__':
    audit = prepare_data()
    print(
        f"Prepared {audit['usable_records']} usable records from {audit['source_records']} workbook records."
    )
